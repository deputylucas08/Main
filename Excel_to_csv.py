import openpyxl
import csv

# Path to your Excel file
excel_file = "C:\\Path\\To\\YourFile.xlsx"

# Load the workbook
wb = openpyxl.load_workbook(excel_file, data_only=True)

# Loop through all sheets in the workbook
for sheet in wb.sheetnames:
    # Get the sheet object
    ws = wb[sheet]
    
    # Create a CSV file for each sheet
    csv_file = f"C:\\Path\\To\\Output\\{sheet}.csv"
    
    # Open the CSV file for writing
    with open(csv_file, mode='w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        
        # Loop through each row in the sheet
        for row in ws.iter_rows(values_only=True):
            # Write the row to the CSV file
            writer.writerow(row)
    
    print(f"Sheet '{sheet}' has been written to {csv_file}")

print("All sheets have been exported to CSV.")
